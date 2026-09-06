"""PHASE-4 — route-level test for the LINE-user directory "Export Report"
button endpoint: GET /admin/api/line-users/export.xlsx.

Uses the real FastAPI app via TestClient with a mocked Supabase client
(_FakeSupabase) seeded exactly like tests/test_line_user_directory.py —
never a real DB or network call. Proves: admin-session auth is enforced,
a real .xlsx is returned with the report filename, the workbook parses,
Thai is intact, and the same search/filter the table uses is honoured.
"""
import io
import re
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from starlette.testclient import TestClient
from openpyxl import load_workbook

from tests.test_business_action_registry import _FakeSupabase
from tests._admin_test_auth import login_as_test_admin
from tests.test_line_user_directory import _seed, _UA, _UB

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class TestLineUsersExportRoute(unittest.TestCase):
    def setUp(self):
        from admin.routes import app
        self.client = TestClient(app)
        self.fake_sb = _seed(_FakeSupabase())
        # export_line_users_xlsx() -> list_line_users() -> _sb() ->
        # services.supabase_client.get_supabase(). Patch that one factory.
        self.patcher = patch("services.supabase_client.get_supabase", return_value=self.fake_sb)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()

    def test_requires_admin_auth(self):
        resp = self.client.get("/admin/api/line-users/export.xlsx", follow_redirects=False)
        self.assertIn(resp.status_code, (302, 307))
        self.assertIn("/admin/login", resp.headers.get("location", ""))

    def test_authenticated_download_is_a_real_xlsx(self):
        login_as_test_admin(self.client)
        resp = self.client.get("/admin/api/line-users/export.xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.headers["content-type"], _XLSX_MIME)
        cd = resp.headers.get("content-disposition", "")
        self.assertRegex(cd, r'filename="line-user-report-\d{8}-\d{4}\.xlsx"')
        # body parses as a workbook, has a header + at least the 3 real users
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "LINE Display Name")
        self.assertGreaterEqual(len(rows) - 1, 3)
        names = {r[0] for r in rows[1:]}
        self.assertIn("สมชาย ใจดี", names)          # Thai intact end-to-end

    def test_export_respects_current_search(self):
        login_as_test_admin(self.client)
        resp = self.client.get("/admin/api/line-users/export.xlsx", params={"search": "anna"})
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        body = list(wb.active.iter_rows(values_only=True))[1:]
        self.assertEqual([r[0] for r in body], ["Anna Wong"])

    def test_no_secret_header_in_export(self):
        login_as_test_admin(self.client)
        resp = self.client.get("/admin/api/line-users/export.xlsx")
        wb = load_workbook(io.BytesIO(resp.content))
        header = " ".join(h or "" for h in list(wb.active.iter_rows(values_only=True))[0]).lower()
        for banned in ("secret", "password", "token", "credential", "apikey", "api_key"):
            self.assertNotIn(banned, header)

    def test_page_has_search_button_and_export_button(self):
        login_as_test_admin(self.client)
        html = self.client.get("/admin/line-users").text
        # explicit "ค้นหา" button that re-runs the same filtered load…
        self.assertIn('id="lu-search-btn"', html)
        self.assertIn('onclick="luLoad()"', html)
        # …placed before the Export Report button in source order
        self.assertLess(html.index('id="lu-search-btn"'), html.index('id="lu-export"'))


if __name__ == "__main__":
    unittest.main()
