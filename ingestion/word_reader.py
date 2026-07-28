from docx import Document
from typing import List, Dict


def read_word(file_path: str) -> str:
    """อ่าน .docx แล้วคืน text ทั้งหมด (backward compat)"""
    pages = read_word_pages(file_path)
    return "\n".join(p["text"] for p in pages).strip()


def read_word_pages(file_path: str) -> List[Dict]:
    """อ่าน .docx คืน list of {page_number, text}
    Word ไม่มี hard page boundary เหมือน PDF — แบ่งทุก ~50 paragraphs แทน"""
    text_parts = []
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
    except Exception as e:
        print(f"❌ อ่าน Word ไม่ได้: {file_path} — {e}")
        return []

    # จัด "หน้า" โดยรวม ~50 paragraph ต่อ 1 logical page
    PARAS_PER_PAGE = 50
    pages = []
    for i in range(0, max(len(text_parts), 1), PARAS_PER_PAGE):
        chunk = "\n".join(text_parts[i:i + PARAS_PER_PAGE])
        if chunk.strip():
            pages.append({
                "page_number": (i // PARAS_PER_PAGE) + 1,
                "text": chunk.strip(),
            })
    return pages
